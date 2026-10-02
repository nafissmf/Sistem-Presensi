"""Presensi Rakernas: import, scan, rekap, dan cetak barcode."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.exceptions import BadRequestException, NotFoundException
from app.core.redis_client import get_redis_client
from app.models.attendance import (
    Attendance,
    AuditLog,
    Event,
    EventParticipant,
    Participant,
)
from app.models.auth import User

JKT = ZoneInfo("Asia/Jakarta")
LIVE_CHANNEL = "rakernas:live"
DEBOUNCE_SECONDS = 3

STATUS_MAP = {
    "aktif": "active",
    "active": "active",
    "terverifikasi": "active",
    "verified": "active",
    "menunggu": "pending",
    "menunggu verifikasi": "pending",
    "pending": "pending",
    "tunda": "hold",
    "tunda bayar": "hold",
    "hold": "hold",
    "ditolak": "rejected",
    "rejected": "rejected",
}


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def now_jkt() -> datetime:
    return datetime.now(JKT)


def publish_live(payload: dict) -> None:
    client = get_redis_client()
    if client is None:
        return
    try:
        client.publish(LIVE_CHANNEL, json.dumps(payload, default=str))
    except Exception:
        return


def write_audit(
    db: Session,
    *,
    user_id: str | None,
    action: str,
    entity: str,
    entity_id: str | None,
    detail: dict | None = None,
) -> None:
    db.add(
        AuditLog(
            user_id=user_id,
            action=action,
            entity=entity,
            entity_id=entity_id,
            detail=json.dumps(detail, default=str) if detail else None,
            created_at=now_utc(),
        )
    )


def normalize_status(raw: str | None) -> str:
    if not raw or not str(raw).strip():
        return "active"
    key = str(raw).strip().lower()
    if key not in STATUS_MAP:
        raise BadRequestException(f"Status tidak dikenal: {raw}")
    return STATUS_MAP[key]


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_participant_file(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            return []
        headers = [_cell(h) for h in data[0]]
        rows = []
        for values in data[1:]:
            rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    else:
        raise BadRequestException("Format file harus .xlsx atau .csv")

    def pick(row: dict, *names: str) -> str:
        lowered = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
        for n in names:
            if n in lowered:
                return _cell(lowered[n])
        return ""

    parsed = []
    for index, row in enumerate(rows, start=2):
        code = pick(row, "kode", "code")
        nama = pick(row, "nama", "name")
        instansi = pick(row, "instansi", "organization", "organisasi")
        status = pick(row, "status")
        if not code and not nama and not instansi:
            continue
        parsed.append(
            {
                "row": index,
                "code": code,
                "name": nama,
                "organization": instansi,
                "status": status,
            }
        )
    return parsed


def preview_import(db: Session, rows: list[dict]) -> dict:
    seen: dict[str, int] = {}
    errors = []
    valid = []
    for row in rows:
        messages = []
        code = row["code"].strip()
        name = row["name"].strip()
        organization = row["organization"].strip()
        if not code:
            messages.append("Kode wajib diisi")
        if not name:
            messages.append("Nama wajib diisi")
        if not organization:
            messages.append("Instansi wajib diisi")
        if code:
            if code in seen:
                messages.append(f"Kode ganda dengan baris {seen[code]}")
            else:
                seen[code] = row["row"]
        status = "active"
        if row["status"]:
            try:
                status = normalize_status(row["status"])
            except BadRequestException as exc:
                messages.append(str(exc.detail))
        if messages:
            errors.append({"row": row["row"], "code": code, "messages": messages})
            continue
        existing = db.query(Participant.id).filter(Participant.code == code).first()
        valid.append(
            {
                "row": row["row"],
                "code": code,
                "name": name,
                "organization": organization,
                "status": status,
                "action": "update" if existing else "create",
            }
        )
    return {
        "total": len(rows),
        "create": sum(1 for r in valid if r["action"] == "create"),
        "update": sum(1 for r in valid if r["action"] == "update"),
        "errors": errors,
        "rows": valid,
    }


def commit_import(db: Session, rows: list[dict], user_id: str | None) -> dict:
    created = 0
    updated = 0
    for row in rows:
        code = row["code"].strip()
        name = row["name"].strip()
        organization = row["organization"].strip()
        status = normalize_status(row.get("status"))
        if not code or not name or not organization:
            raise BadRequestException("Baris import tidak lengkap")
        existing = db.query(Participant).filter(Participant.code == code).first()
        if existing:
            existing.name = name
            existing.organization = organization
            existing.status = status
            updated += 1
        else:
            db.add(Participant(code=code, name=name, organization=organization, status=status))
            created += 1
    write_audit(
        db,
        user_id=user_id,
        action="import",
        entity="participant",
        entity_id=None,
        detail={"created": created, "updated": updated},
    )
    db.commit()
    return {"created": created, "updated": updated}


def expected_count(db: Session, event: Event) -> int:
    if event.audience == "subset":
        return (
            db.query(func.count(EventParticipant.participant_id))
            .filter(EventParticipant.event_id == event.id)
            .scalar()
            or 0
        )
    return (
        db.query(func.count(Participant.id)).filter(Participant.status != "rejected").scalar() or 0
    )


def present_count(db: Session, event_id: str) -> int:
    return db.query(func.count(Attendance.id)).filter(Attendance.event_id == event_id).scalar() or 0


def counts(db: Session, event: Event) -> dict:
    present = present_count(db, event.id)
    expected = expected_count(db, event)
    return {
        "present": present,
        "expected": expected,
        "percent": round((present / expected) * 100, 1) if expected else 0,
    }


def _window_message(event: Event, moment: datetime) -> str | None:
    if event.status == "draft":
        return "Kegiatan belum dibuka"
    if event.status in {"done", "archived"}:
        return "Kegiatan sudah ditutup"
    if event.status != "active":
        return "Kegiatan belum dibuka"
    if event.scan_open_before_min is None and event.scan_close_after_min is None:
        return None
    start = datetime.combine(event.event_date, event.start_time, tzinfo=JKT)
    end = datetime.combine(event.event_date, event.end_time, tzinfo=JKT)
    open_at = start - timedelta(minutes=event.scan_open_before_min or 0)
    close_at = end + timedelta(minutes=event.scan_close_after_min or 0)
    if moment < open_at:
        return "Kegiatan belum dibuka"
    if moment > close_at:
        return "Kegiatan sudah ditutup"
    return None


def _participant_payload(participant: Participant) -> dict:
    return {
        "id": participant.id,
        "code": participant.code,
        "name": participant.name,
        "organization": participant.organization,
        "status": participant.status,
    }


def scan_code(
    db: Session,
    *,
    event_id: str,
    code: str,
    method: str,
    user: User,
    device: str | None,
    force: bool = False,
) -> dict:
    event = db.query(Event).filter(Event.id == event_id).first()
    if not event:
        raise NotFoundException("Kegiatan tidak ditemukan")

    cleaned = code.strip()
    if not cleaned:
        raise BadRequestException("Kode kosong")

    client = get_redis_client()
    debounce_key = f"rakernas:debounce:{event_id}:{cleaned}"
    if client is not None:
        try:
            if not client.set(debounce_key, "1", nx=True, ex=DEBOUNCE_SECONDS):
                existing = (
                    db.query(Attendance)
                    .join(Participant)
                    .filter(Attendance.event_id == event_id, Participant.code == cleaned)
                    .first()
                )
                participant = db.query(Participant).filter(Participant.code == cleaned).first()
                return {
                    "result": "debounced",
                    "message": "Scan yang sama diabaikan",
                    "participant": _participant_payload(participant) if participant else None,
                    "scanned_at": existing.scanned_at.isoformat() if existing else None,
                    "counts": counts(db, event),
                }
        except Exception:
            pass

    closed = _window_message(event, now_jkt())
    participant = db.query(Participant).filter(Participant.code == cleaned).first()
    if closed:
        return {
            "result": "closed",
            "message": closed,
            "participant": _participant_payload(participant) if participant else None,
            "scanned_at": None,
            "counts": counts(db, event),
        }
    if not participant:
        matches = (
            db.query(Participant)
            .filter(Participant.name.ilike(f"%{cleaned}%"))
            .limit(5)
            .all()
        )
        if len(matches) == 1:
            participant = matches[0]
        elif len(matches) > 1:
            return {
                "result": "unknown",
                "message": "Nama cocok dengan lebih dari satu peserta, pakai kode",
                "participant": None,
                "scanned_at": None,
                "counts": counts(db, event),
            }
        else:
            return {
                "result": "unknown",
                "message": "Kode tidak ditemukan",
                "participant": None,
                "scanned_at": None,
                "counts": counts(db, event),
            }
    if participant.status == "rejected":
        return {
            "result": "rejected",
            "message": "Peserta ditolak",
            "participant": _participant_payload(participant),
            "scanned_at": None,
            "counts": counts(db, event),
        }

    if event.audience == "subset":
        invited = (
            db.query(EventParticipant)
            .filter(
                EventParticipant.event_id == event.id,
                EventParticipant.participant_id == participant.id,
            )
            .first()
        )
        if not invited and not force:
            return {
                "result": "not_invited",
                "message": "Bukan peserta kegiatan ini",
                "participant": _participant_payload(participant),
                "scanned_at": None,
                "counts": counts(db, event),
            }

    existing = (
        db.query(Attendance)
        .filter(Attendance.event_id == event.id, Attendance.participant_id == participant.id)
        .first()
    )
    if existing and not event.allow_duplicate:
        local = existing.scanned_at.astimezone(JKT)
        return {
            "result": "already",
            "message": f"Sudah hadir pukul {local.strftime('%H.%M')}",
            "participant": _participant_payload(participant),
            "scanned_at": existing.scanned_at.isoformat(),
            "attendance_id": existing.id,
            "counts": counts(db, event),
        }

    scanned_at = now_utc()
    if existing and event.allow_duplicate:
        existing.scanned_at = scanned_at
        existing.scanned_by = user.id
        existing.device = (device or "")[:255] or None
        existing.method = method
        attendance = existing
        result = "success"
        message = "Kehadiran diperbarui"
    else:
        attendance = Attendance(
            event_id=event.id,
            participant_id=participant.id,
            scanned_at=scanned_at,
            scanned_by=user.id,
            device=(device or "")[:255] or None,
            method=method,
        )
        db.add(attendance)
        db.flush()
        result = "success"
        message = "Hadir"

    try:
        write_audit(
            db,
            user_id=user.id,
            action="scan",
            entity="attendance",
            entity_id=attendance.id,
            detail={"code": participant.code, "event_id": event.id, "method": method, "result": result},
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        existing = (
            db.query(Attendance)
            .filter(Attendance.event_id == event.id, Attendance.participant_id == participant.id)
            .first()
        )
        local = existing.scanned_at.astimezone(JKT) if existing else now_jkt()
        return {
            "result": "already",
            "message": f"Sudah hadir pukul {local.strftime('%H.%M')}",
            "participant": _participant_payload(participant),
            "scanned_at": existing.scanned_at.isoformat() if existing else None,
            "counts": counts(db, event),
        }

    db.refresh(attendance)
    payload = {
        "type": "scan",
        "result": result,
        "event_id": event.id,
        "attendance": {
            "id": attendance.id,
            "scanned_at": attendance.scanned_at.isoformat(),
            "method": attendance.method,
            "participant": _participant_payload(participant),
            "event_name": event.name,
        },
        "counts": counts(db, event),
    }
    publish_live(payload)
    return {
        "result": result,
        "message": message,
        "participant": _participant_payload(participant),
        "scanned_at": attendance.scanned_at.isoformat(),
        "attendance_id": attendance.id,
        "counts": payload["counts"],
    }


def undo_last_scan(db: Session, *, event_id: str, user: User) -> dict:
    event = db.query(Event).filter(Event.id == event_id).first()
    if not event:
        raise NotFoundException("Kegiatan tidak ditemukan")
    last = (
        db.query(Attendance)
        .filter(Attendance.event_id == event_id, Attendance.scanned_by == user.id)
        .order_by(Attendance.scanned_at.desc())
        .first()
    )
    if not last:
        raise NotFoundException("Tidak ada scan yang bisa dibatalkan")
    participant = last.participant
    db.delete(last)
    write_audit(
        db,
        user_id=user.id,
        action="undo_scan",
        entity="attendance",
        entity_id=last.id,
        detail={"code": participant.code if participant else None, "event_id": event_id},
    )
    db.commit()
    publish_live({"type": "undo", "event_id": event_id, "counts": counts(db, event)})
    return {"message": "Scan terakhir dibatalkan", "counts": counts(db, event)}


def recent_scans(db: Session, event_id: str, limit: int = 10, user_id: str | None = None) -> list[dict]:
    query = db.query(Attendance).filter(Attendance.event_id == event_id)
    if user_id:
        query = query.filter(Attendance.scanned_by == user_id)
    rows = query.order_by(Attendance.scanned_at.desc()).limit(limit).all()
    items = []
    for row in rows:
        local = row.scanned_at.astimezone(JKT)
        items.append(
            {
                "id": row.id,
                "scanned_at": row.scanned_at.isoformat(),
                "time_label": local.strftime("%H.%M"),
                "method": row.method,
                "participant": _participant_payload(row.participant) if row.participant else None,
            }
        )
    return items


def suggest_event(db: Session) -> Event | None:
    moment = now_jkt()
    active = (
        db.query(Event)
        .filter(Event.status == "active", Event.event_date == moment.date())
        .order_by(Event.start_time)
        .all()
    )
    for event in active:
        if _window_message(event, moment) is None:
            start = datetime.combine(event.event_date, event.start_time, tzinfo=JKT)
            end = datetime.combine(event.event_date, event.end_time, tzinfo=JKT)
            if start <= moment <= end:
                return event
    return active[0] if active else None


def dashboard_snapshot(db: Session) -> dict:
    today = now_jkt().date()
    total = db.query(func.count(Participant.id)).filter(Participant.status != "rejected").scalar() or 0
    today_start = datetime.combine(today, time.min, tzinfo=JKT).astimezone(timezone.utc)
    today_end = datetime.combine(today, time.max, tzinfo=JKT).astimezone(timezone.utc)
    present_today = (
        db.query(func.count(func.distinct(Attendance.participant_id)))
        .filter(Attendance.scanned_at >= today_start, Attendance.scanned_at <= today_end)
        .scalar()
        or 0
    )
    events = (
        db.query(Event)
        .filter(Event.event_date == today, Event.status != "archived")
        .order_by(Event.start_time)
        .all()
    )
    current = suggest_event(db)
    event_cards = []
    for event in events:
        card = {
            "id": event.id,
            "name": event.name,
            "status": event.status,
            "location": event.location,
            "color": event.color,
            "start_time": event.start_time.strftime("%H:%M"),
            "end_time": event.end_time.strftime("%H:%M"),
            **counts(db, event),
        }
        event_cards.append(card)

    active_counts = counts(db, current) if current else {"present": 0, "expected": total, "percent": 0}
    feed_rows = (
        db.query(Attendance)
        .filter(Attendance.scanned_at >= today_start)
        .order_by(Attendance.scanned_at.desc())
        .limit(12)
        .all()
    )
    feed = []
    for row in feed_rows:
        local = row.scanned_at.astimezone(JKT)
        feed.append(
            {
                "id": row.id,
                "name": row.participant.name if row.participant else "",
                "organization": row.participant.organization if row.participant else "",
                "event_name": row.event.name if row.event else "",
                "time_label": local.strftime("%H.%M"),
            }
        )

    hourly = [0] * 24
    day_rows = (
        db.query(Attendance.scanned_at)
        .filter(Attendance.scanned_at >= today_start, Attendance.scanned_at <= today_end)
        .all()
    )
    for (scanned_at,) in day_rows:
        hourly[scanned_at.astimezone(JKT).hour] += 1

    org_rows = (
        db.query(Participant.organization, func.count(Attendance.id))
        .join(Attendance, Attendance.participant_id == Participant.id)
        .filter(Attendance.scanned_at >= today_start, Attendance.scanned_at <= today_end)
        .group_by(Participant.organization)
        .order_by(func.count(Attendance.id).desc())
        .limit(8)
        .all()
    )
    return {
        "total_participants": total,
        "present_today": present_today,
        "active_event": (
            {"id": current.id, "name": current.name, **active_counts} if current else None
        ),
        "events": event_cards,
        "feed": feed,
        "hourly": [{"hour": f"{h:02d}", "count": hourly[h]} for h in range(24)],
        "organizations": [{"name": name, "count": count} for name, count in org_rows],
    }


def event_report_rows(db: Session, event_id: str) -> tuple[Event, list[dict]]:
    event = db.query(Event).filter(Event.id == event_id).first()
    if not event:
        raise NotFoundException("Kegiatan tidak ditemukan")
    if event.audience == "subset":
        participants = (
            db.query(Participant)
            .join(EventParticipant, EventParticipant.participant_id == Participant.id)
            .filter(EventParticipant.event_id == event.id)
            .order_by(Participant.name)
            .all()
        )
    else:
        participants = (
            db.query(Participant)
            .filter(Participant.status != "rejected")
            .order_by(Participant.name)
            .all()
        )
    attendance_map = {
        row.participant_id: row
        for row in db.query(Attendance).filter(Attendance.event_id == event.id).all()
    }
    rows = []
    for participant in participants:
        attendance = attendance_map.get(participant.id)
        local = attendance.scanned_at.astimezone(JKT) if attendance else None
        rows.append(
            {
                "code": participant.code,
                "name": participant.name,
                "organization": participant.organization,
                "status": "Hadir" if attendance else "Tidak hadir",
                "scanned_at": local.strftime("%Y-%m-%d %H:%M") if local else "",
            }
        )
    return event, rows


def report_workbook(db: Session, event_id: str) -> tuple[str, bytes]:
    event, rows = event_report_rows(db, event_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kehadiran"
    ws.append(["Kode", "Nama", "Instansi", "Status", "Jam hadir"])
    for row in rows:
        ws.append([row["code"], row["name"], row["organization"], row["status"], row["scanned_at"]])
    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"kehadiran-{event.name}".replace(" ", "-") + ".xlsx"
    return filename, buffer.getvalue()


def _wrap_text(text: str, max_chars: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if len(trial) <= max_chars:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def barcode_pdf(db: Session, participant_ids: list[str], symbology: str, user_id: str | None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdf_canvas

    participants = (
        db.query(Participant)
        .filter(Participant.id.in_(participant_ids))
        .order_by(Participant.code)
        .all()
    )
    if not participants:
        raise NotFoundException("Peserta tidak ditemukan")

    buffer = io.BytesIO()
    pdf = pdf_canvas.Canvas(buffer, pagesize=A4)
    page_w, page_h = A4
    card_w, card_h = 85 * mm, 55 * mm
    margin_x, margin_y, gap = 12 * mm, 12 * mm, 4 * mm
    cols = int((page_w - 2 * margin_x + gap) // (card_w + gap))
    rows = int((page_h - 2 * margin_y + gap) // (card_h + gap))
    per_page = max(cols * rows, 1)

    for index, participant in enumerate(participants):
        slot = index % per_page
        if index > 0 and slot == 0:
            pdf.showPage()
        col = slot % cols
        row = slot // cols
        x = margin_x + col * (card_w + gap)
        y = page_h - margin_y - (row + 1) * card_h - row * gap
        pdf.setStrokeColorRGB(0.8, 0.82, 0.86)
        pdf.roundRect(x, y, card_w, card_h, 3 * mm, stroke=1, fill=0)
        pdf.setFillColorRGB(0.07, 0.09, 0.15)
        name_lines = _wrap_text(participant.name, 32)[:2]
        org_lines = _wrap_text(participant.organization, 38)[:2]
        pdf.setFont("Helvetica-Bold", 9 if len(participant.name) > 32 else 11)
        text_y = y + card_h - 8 * mm
        for line in name_lines:
            pdf.drawString(x + 4 * mm, text_y, line[:48])
            text_y -= 4 * mm
        pdf.setFont("Helvetica", 7 if len(participant.organization) > 40 else 8)
        pdf.setFillColorRGB(0.28, 0.33, 0.4)
        for line in org_lines:
            pdf.drawString(x + 4 * mm, text_y, line[:60])
            text_y -= 3.4 * mm

        image = _barcode_image(participant.code, symbology)
        img_w, img_h = 74 * mm, 16 * mm
        pdf.drawImage(ImageReader(image), x + 5 * mm, y + 8 * mm, width=img_w, height=img_h, mask="auto")
        pdf.setFillColorRGB(0.07, 0.09, 0.15)
        pdf.setFont("Courier", 8)
        pdf.drawCentredString(x + card_w / 2, y + 4 * mm, participant.code)
        participant.printed_at = now_utc()

    write_audit(
        db,
        user_id=user_id,
        action="print_barcode",
        entity="participant",
        entity_id=None,
        detail={"count": len(participants), "symbology": symbology},
    )
    db.commit()
    pdf.save()
    return buffer.getvalue()


def _barcode_image(code: str, symbology: str):
    import io as _io

    from PIL import Image

    if symbology == "qr":
        import qrcode

        img = qrcode.make(code)
        out = _io.BytesIO()
        img.save(out, format="PNG")
        out.seek(0)
        return out

    from barcode import Code128
    from barcode.writer import ImageWriter

    writer = ImageWriter()
    buffer = _io.BytesIO()
    Code128(code, writer=writer).write(
        buffer,
        options={
            "module_width": 0.28,
            "module_height": 10,
            "quiet_zone": 2,
            "write_text": False,
            "font_size": 1,
        },
    )
    buffer.seek(0)
    image = Image.open(buffer).convert("RGB")
    out = _io.BytesIO()
    image.save(out, format="PNG")
    out.seek(0)
    return out


def resolve_participant_ids(db: Session, raw_ids: list[str]) -> list[str]:
    if not raw_ids:
        return []
    found_ids = {row[0] for row in db.query(Participant.id).filter(Participant.id.in_(raw_ids)).all()}
    missing = [item for item in raw_ids if item not in found_ids]
    code_map = {}
    if missing:
        for participant in db.query(Participant).filter(Participant.code.in_(missing)).all():
            code_map[participant.code] = participant.id
    unknown = [item for item in missing if item not in code_map]
    if unknown:
        raise BadRequestException("Peserta tidak ditemukan: " + ", ".join(unknown[:8]))
    return [item if item in found_ids else code_map[item] for item in raw_ids]


def replace_event_participants(db: Session, event: Event, participant_ids: list[str]) -> None:
    db.query(EventParticipant).filter(EventParticipant.event_id == event.id).delete()
    if event.audience != "subset":
        return
    unique_ids = list(dict.fromkeys(participant_ids))
    found = db.query(Participant.id).filter(Participant.id.in_(unique_ids)).all()
    if len(found) != len(unique_ids):
        raise BadRequestException("Ada peserta yang tidak ditemukan")
    for participant_id in unique_ids:
        db.add(EventParticipant(event_id=event.id, participant_id=participant_id))


def duplicate_event(db: Session, event: Event, user_id: str | None) -> Event:
    copy = Event(
        name=f"{event.name} (salinan)",
        event_date=event.event_date,
        start_time=event.start_time,
        end_time=event.end_time,
        location=event.location,
        status="draft",
        audience=event.audience,
        scan_open_before_min=event.scan_open_before_min,
        scan_close_after_min=event.scan_close_after_min,
        allow_duplicate=event.allow_duplicate,
        color=event.color,
    )
    db.add(copy)
    db.flush()
    if event.audience == "subset":
        for link in event.event_participants:
            db.add(EventParticipant(event_id=copy.id, participant_id=link.participant_id))
    write_audit(db, user_id=user_id, action="duplicate", entity="event", entity_id=copy.id, detail={"from": event.id})
    db.commit()
    db.refresh(copy)
    return copy
